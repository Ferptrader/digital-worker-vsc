"""Fail-closed, evidence-bound document extraction. No regulatory approval logic."""
from __future__ import annotations
from pathlib import Path
import hashlib
import json

ALLOWED_TYPES = {'.txt', '.md', '.pdf', '.docx'}
ANALYSIS_TYPES = {'requirements', 'risk', 'compliance', 'structure'}
MAX_BYTES = 5_000_000


def analyze_document(document_path: str, analysis_type: str, allowed_root: str) -> dict:
    result = {'status':'NOT_VERIFIED', 'analysis_type':analysis_type,
              'gxp_compliance':'NOT_VERIFIED', 'gxp_approval':'NOT_GRANTED',
              'gxp_release':'NOT_GRANTED', 'evidence':None, 'reason':None}
    try:
        if analysis_type not in ANALYSIS_TYPES:
            raise ValueError('INVALID_ANALYSIS_TYPE')
        root = Path(allowed_root).resolve(strict=True)
        path = Path(document_path).resolve(strict=True)
        if not path.is_relative_to(root) or path == root:
            raise ValueError('PATH_OUTSIDE_SCOPE')
        if not path.is_file() or path.suffix.lower() not in ALLOWED_TYPES:
            raise ValueError('UNSUPPORTED_DOCUMENT')
        if path.stat().st_size > MAX_BYTES:
            raise ValueError('DOCUMENT_TOO_LARGE')
        content_bytes = path.read_bytes()
        if not content_bytes:
            raise ValueError('EMPTY_DOCUMENT')
        if len(content_bytes) > MAX_BYTES:
            raise ValueError('DOCUMENT_TOO_LARGE')
        doc_sha = hashlib.sha256(content_bytes).hexdigest()
        suffix = path.suffix.lower()
        if suffix in ('.txt', '.md'):
            text = content_bytes.decode('utf-8')
        elif suffix == '.pdf':
            from io import BytesIO
            from pypdf import PdfReader
            pages = PdfReader(BytesIO(content_bytes), strict=True).pages
            text = '\n'.join(p.extract_text() or '' for p in pages)
            # PDF flattening may erase revision semantics: extraction is never certified as active text.
        else:
            from io import BytesIO
            from docx import Document
            d = Document(BytesIO(content_bytes))
            # Native DOCX revision markup is not reliably visible via paragraph.text.
            # Refuse unassessed redlines, preventing deleted terms being treated as active.
            from zipfile import ZipFile
            with ZipFile(BytesIO(content_bytes)) as z:
                xml = z.read('word/document.xml')
            from xml.etree import ElementTree as ET
            # Inspect WordprocessingML elements independently of XML namespace prefixes.
            review_tags = {'del','ins','delText','moveFrom','moveTo','strike','dstrike','trackRevisions'}
            parsed = ET.fromstring(xml)
            if any(e.tag.rsplit('}',1)[-1] in review_tags for e in parsed.iter()):
                raise ValueError('DOCX_REVISIONS_NEED_DEDICATED_REVIEW')
            text = '\n'.join([p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells])
        if not text.strip():
            raise ValueError('NO_EXTRACTABLE_TEXT')
        # Extraction is NOT an analytical or regulatory conclusion.
        result['status'] = 'EXTRACTED_UNVERIFIED'
        result['evidence'] = {'sha256':doc_sha,'format':suffix,'characters':len(text),
                              'excerpt':text[:2000]}
        result['reason'] = 'EXTRACTION_ONLY_REQUIRES_HUMAN_REVIEW'
    except Exception as exc:  # Boundary is fail-closed; never surface parser exception as PASS
        result['reason'] = str(exc) if isinstance(exc, ValueError) else type(exc).__name__
    return result


def analyze_json(document_path: str, analysis_type: str, allowed_root: str) -> str:
    return json.dumps(analyze_document(document_path, analysis_type, allowed_root), ensure_ascii=False, sort_keys=True)
