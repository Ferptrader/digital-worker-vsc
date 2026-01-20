"""
Digital Worker VSC - Knowledge Base Document Reader
Lê e processa documentos de referência (PDF, Word, Excel)
"""

from pathlib import Path
import PyPDF2
from docx import Document
import openpyxl
import json

class DocumentReader:
    def __init__(self, kb_path="knowledge_base"):
        self.kb_path = Path(kb_path)
        self.categories = ['manuais', 'especificacoes', 'documentos_empresa', 'normas', 'projeto_atual']
        self.knowledge = {}
    
    def ensure_folders(self):
        """Cria estrutura de pastas se não existir"""
        for category in self.categories:
            folder = self.kb_path / category
            folder.mkdir(parents=True, exist_ok=True)
    
    def scan_all_documents(self):
        """Escaneia todos os documentos"""
        print("\n" + "="*60)
        print("🔍 ESCANEANDO KNOWLEDGE BASE")
        print("="*60)
        
        docs = {cat: [] for cat in self.categories}
        
        for category in self.categories:
            folder = self.kb_path / category
            if folder.exists():
                docs[category].extend(list(folder.glob("*.pdf")))
                docs[category].extend(list(folder.glob("*.docx")))
                docs[category].extend(list(folder.glob("*.xlsx")))
        
        total = sum(len(v) for v in docs.values())
        print(f"\n📊 TOTAL: {total} documentos encontrados")
        
        for cat, files in docs.items():
            if files:
                print(f"\n📁 {cat.upper()}: {len(files)} arquivo(s)")
                for f in files:
                    print(f"   → {f.name}")
        
        return docs
    
    def read_pdf(self, file_path):
        """Extrai texto de PDF"""
        try:
            text = ""
            with open(file_path, 'rb') as file:
                pdf = PyPDF2.PdfReader(file)
                for page in pdf.pages:
                    text += page.extract_text() + "\n"
            return text
        except Exception as e:
            return f"Erro ao ler PDF: {e}"
    
    def read_word(self, file_path):
        """Extrai texto de Word"""
        try:
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
            return text
        except Exception as e:
            return f"Erro ao ler Word: {e}"
    
    def read_excel(self, file_path):
        """Extrai dados de Excel"""
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = ""
            for sheet in wb.worksheets:
                text += f"\n=== Planilha: {sheet.title} ===\n"
                for row in sheet.iter_rows(values_only=True, max_row=100):
                    text += " | ".join([str(c) for c in row if c]) + "\n"
            return text
        except Exception as e:
            return f"Erro ao ler Excel: {e}"
    
    def extract_all_content(self):
        """Extrai conteúdo de TODOS os documentos"""
        all_docs = self.scan_all_documents()
        self.knowledge = {}
        
        for category, files in all_docs.items():
            self.knowledge[category] = {}
            
            for file_path in files:
                print(f"\n📖 Processando: {file_path.name}...")
                
                try:
                    if file_path.suffix == '.pdf':
                        content = self.read_pdf(file_path)
                    elif file_path.suffix == '.docx':
                        content = self.read_word(file_path)
                    elif file_path.suffix == '.xlsx':
                        content = self.read_excel(file_path)
                    else:
                        content = ""
                    
                    self.knowledge[category][file_path.name] = {
                        'path': str(file_path),
                        'content': content[:10000],
                        'full_size': len(content),
                        'type': file_path.suffix
                    }
                    print(f"   ✅ {len(content)} caracteres extraídos")
                    
                except Exception as e:
                    print(f"   ❌ Erro: {e}")
        
        return self.knowledge
    
    def get_summary(self):
        """Gera resumo da Knowledge Base"""
        if not self.knowledge:
            self.extract_all_content()
        
        summary = "\n📚 KNOWLEDGE BASE CARREGADA\n" + "="*60 + "\n"
        
        for category, docs in self.knowledge.items():
            if docs:
                summary += f"\n{category.upper()}: {len(docs)} documento(s)\n"
                for filename in docs.keys():
                    summary += f"  • {filename}\n"
        
        total_docs = sum(len(d) for d in self.knowledge.values())
        summary += f"\n{'='*60}\nTOTAL: {total_docs} documentos processados\n"
        
        return summary
    
    def search(self, query, max_results=5):
        """Busca em todos os documentos"""
        results = []
        
        for category, docs in self.knowledge.items():
            for filename, data in docs.items():
                if query.lower() in data['content'].lower():
                    idx = data['content'].lower().find(query.lower())
                    context = data['content'][max(0, idx-300):idx+300]
                    
                    results.append({
                        'file': filename,
                        'category': category,
                        'context': context,
                        'relevance': 'high'
                    })
        
        return results[:max_results]

if __name__ == "__main__":
    reader = DocumentReader()
    reader.ensure_folders()
    knowledge = reader.extract_all_content()
    print(reader.get_summary())
